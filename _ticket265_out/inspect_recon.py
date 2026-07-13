"""Analyze a reconciled xlsx for #265 offenders: flux tagged 'libre' whose
written interval is collapsed (0 <= width <= 1e-3 rel) or inverted (sup < inf).
Also dumps the 'Options de reconciliation' sheet.  Usage: inspect_recon.py <xlsx>
"""
import math
import sys
import openpyxl

path = sys.argv[1]
wb = openpyxl.load_workbook(path, data_only=True, read_only=True)

# --- reconciliation options actually used ---
if "Options de réconciliation" in wb.sheetnames:
    print("=== Options de réconciliation ===")
    for row in wb["Options de réconciliation"].iter_rows(values_only=True):
        cells = [str(c) for c in row if c is not None and str(c).strip() != ""]
        if cells:
            print("  " + " | ".join(cells))

ws = wb["Analyses des résultats"]
rows = list(ws.iter_rows(values_only=True))
header = [str(c) if c is not None else "" for c in rows[0]]
idx = {h: i for i, h in enumerate(header)}
ci_t, ci_i, ci_s = idx["Type de variable"], idx["Borne inférieure"], idx["Borne supérieure"]
ci_v, ci_o, ci_d, ci_y = idx["Valeur reconciliée"], idx["Origine"], idx["Destination"], idx["Année"]


def num(x):
    try:
        return float(x)
    except (TypeError, ValueError):
        return None


n_libre = 0
widths = []
offenders = []
for r in rows[1:]:
    t = str(r[ci_t]) if r[ci_t] is not None else ""
    if "libre" not in t:
        continue
    n_libre += 1
    inf, sup, val = num(r[ci_i]), num(r[ci_s]), num(r[ci_v])
    if inf is None or sup is None:
        continue
    scale = max(abs(inf), abs(sup), 1e-12)
    widths.append((sup - inf) / scale)
    inverted = sup < inf
    collapsed = 0 <= (sup - inf) <= 1e-3 * scale
    if inverted or collapsed:
        offenders.append((r[ci_o], r[ci_d], r[ci_y], val, inf, sup,
                          "inverted" if inverted else "collapsed"))

print(f"\n=== #265 analysis on: {path.split(chr(92))[-1]} ===")
print("rows with Type de variable containing 'libre':", n_libre)
n_inv = sum(1 for o in offenders if o[6] == "inverted")
print(f"OFFENDERS (collapsed or inverted): {len(offenders)}"
      f"  [collapsed={len(offenders) - n_inv}, inverted={n_inv}]")

buckets = {}
for w in widths:
    aw = abs(w)
    b = -99 if aw == 0 else int(math.floor(math.log10(aw)))
    buckets[b] = buckets.get(b, 0) + 1
print("log10(|signed rel width|) histogram for 'libre':")
for b in sorted(buckets):
    label = "exact 0" if b == -99 else f"1e{b:+d}"
    print(f"   {label:>8} : {buckets[b]}")

print("sample offenders:")
for o in offenders[:6]:
    print(f"   {o[0]} -> {o[1]} [{o[2]}] val={o[3]:.4g} "
          f"inf={o[4]:.6g} sup={o[5]:.6g} ({o[6]})")
